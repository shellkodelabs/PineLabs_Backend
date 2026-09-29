"""
Parses an uploaded Bin Series spreadsheet (.xlsx or .csv/.txt) into the
SAME row schema (GiftCardBulkUploadRowRequest / WalletBulkUploadRowRequest)
the existing bulk-upload business logic already accepts
(app.services.gift_card_bin_service.bulk_upload_gift_card_bin_records /
app.services.wallet_bin_service.bulk_upload_wallet_bin_records). This
module does ONLY file-format parsing and header-to-column mapping — no
DB access, no business-rule validation (required fields, uniqueness,
issuer matching). Every one of those checks already exists, unchanged,
in the two per-type bulk-upload services; a bad/missing/malformed CELL
value is left for that existing per-row validation to report exactly as
it already does for a directly-supplied JSON body. This module only
rejects the FILE itself (unreadable, empty, no detectable header row,
too many rows).

Added as a follow-up to Stage 6, wiring the Gift Card/Wallet
bulk-upload routers to accept a real uploaded file (multipart/
form-data, FastAPI `UploadFile`) instead of requiring the caller to
have already parsed it into JSON. Confirmed against the actual project
state before writing this: the approved frontend (PineLabs_Frontend's
UploadSheetModal.jsx/utils/csv.js) parses the file entirely client-side
today and does not yet call the backend for bulk upload at all — there
is no existing frontend FormData field name or wire format to match.
This module's header-detection approach mirrors, independently
(duplicated, not imported — different language/runtime, no shared
module boundary), the SAME strategy utils/csv.js's own
headerToColumn/HEADER_HINTS uses: an exact, normalized (lowercased,
punctuation/whitespace-stripped) match against each column's own
key/label first, then a small fuzzy-hint table for common header
spellings ("BIN / IIN", "Merchant Pfx", "Program Type", etc.) — so a
spreadsheet the frontend would successfully read locally also parses
correctly here.

SUPPORTED FILE TYPES: .xlsx (via openpyxl) and .csv/.txt (via the
stdlib `csv` module, decoded as UTF-8 with a BOM stripped if present).
Legacy binary .xls is NOT supported (openpyxl cannot read it, and no
extra library is being added just for that) — rejected with a clear
error naming the unsupported extension.
"""
import csv
import io
import re
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from app.core.exceptions import ValidationError
from app.schemas.gift_card_bin_series import GiftCardBulkUploadRowRequest
from app.schemas.wallet_bin_series import WalletBulkUploadRowRequest

_MAX_ROWS = 1000  # matches GiftCardBulkUploadRequest/WalletBulkUploadRequest's rows max_length.

# (column key, canonical label, fuzzy-hint substrings — tried longest
# first across ALL columns, so a more specific hint like "merchantprefix"
# is never shadowed by a shorter, more generic one like "merchant").
_GIFT_CARD_COLUMNS: List[Tuple[str, str, List[str]]] = [
    ("instance", "Instance", ["instance"]),
    ("issuer", "Issuer", ["issuer"]),
    ("merchant", "Merchant", ["merchant"]),
    ("cardProgramGroupName", "CardProgramGroupName", ["cardprogramgroupname", "programgroupname"]),
    ("binIin", "BIN", ["biniin", "bin", "iin"]),
    ("merchantPrefix", "Merchant Prefix", ["merchantprefix", "merchantpfx", "prefix"]),
    ("cardProgramGroupType", "CardProgramGroupType", ["cardprogramgrouptype", "programgrouptype", "programtype"]),
    ("cardType", "CardType", ["cardtype"]),
    ("ticketNumber", "Ticket Number", ["ticketnumber", "ticket"]),
]

_WALLET_COLUMNS: List[Tuple[str, str, List[str]]] = [
    ("instance", "Instance", ["instance"]),
    ("issuer", "Issuer", ["issuer"]),
    ("merchant", "Merchant", ["merchant"]),
    ("walletProgramName", "Wallet Program Name", ["walletprogramname", "programname"]),
    ("binIin", "BIN", ["biniin", "bin", "iin"]),
    ("merchantPrefix", "Merchant Prefix", ["merchantprefix", "merchantpfx", "prefix"]),
    ("walletProgramGroupType", "Wallet Program Group Type", ["walletprogramgrouptype", "programgrouptype", "programtype"]),
    ("ticketNumber", "Ticket Number", ["ticketnumber", "ticket"]),
]


def _normalize(text: Optional[str]) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def _read_table(filename: str, content: bytes) -> List[List[str]]:
    """Returns a 2D list of string cells (row 0 = header row). Raises
    ValidationError for an unsupported/unreadable/empty file."""
    lower_name = (filename or "").lower()

    if lower_name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError(
                "Could not read the uploaded file as a valid .xlsx workbook.",
                code="BIN_UPLOAD_FILE_UNREADABLE",
            ) from exc
        sheet = workbook.worksheets[0]
        table = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif lower_name.endswith(".xls"):
        raise ValidationError(
            "Legacy .xls files are not supported — please save the file as .xlsx or .csv and upload again.",
            code="BIN_UPLOAD_FILE_UNSUPPORTED_TYPE",
        )
    else:
        # .csv / .txt / no recognized extension — treat as delimited text.
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValidationError(
                "Could not read the uploaded file as UTF-8 text CSV.", code="BIN_UPLOAD_FILE_UNREADABLE"
            ) from exc
        table = list(csv.reader(io.StringIO(text)))

    table = [row for row in table if any((cell or "").strip() for cell in row)]
    if not table:
        raise ValidationError("The uploaded file is empty.", code="BIN_UPLOAD_FILE_EMPTY")
    return table


def _build_header_map(header_row: List[str], columns: List[Tuple[str, str, List[str]]]) -> Dict[int, str]:
    exact_lookup: Dict[str, str] = {}
    hints: List[Tuple[str, str]] = []
    for key, label, hint_list in columns:
        exact_lookup[_normalize(key)] = key
        exact_lookup[_normalize(label)] = key
        for hint in hint_list:
            hints.append((hint, key))
    hints.sort(key=lambda pair: len(pair[0]), reverse=True)

    column_map: Dict[int, str] = {}
    assigned_keys = set()
    for index, cell in enumerate(header_row):
        normalized = _normalize(cell)
        if not normalized:
            continue
        matched_key = exact_lookup.get(normalized)
        if matched_key is None:
            for hint, key in hints:
                if hint in normalized:
                    matched_key = key
                    break
        if matched_key is not None and matched_key not in assigned_keys:
            column_map[index] = matched_key
            assigned_keys.add(matched_key)

    return column_map


def _rows_from_table(table: List[List[str]], columns: List[Tuple[str, str, List[str]]]) -> List[Dict[str, Optional[str]]]:
    header_row, *data_rows = table
    column_map = _build_header_map(header_row, columns)
    if not column_map:
        raise ValidationError(
            "Could not detect any recognizable Bin Series column headers in the uploaded file's first row.",
            code="BIN_UPLOAD_HEADERS_NOT_DETECTED",
        )

    parsed_rows: List[Dict[str, Optional[str]]] = []
    for row in data_rows:
        if not any((cell or "").strip() for cell in row):
            continue  # blank row — skipped, never a failure.
        row_dict: Dict[str, Optional[str]] = {}
        for index, key in column_map.items():
            raw_value = row[index] if index < len(row) else ""
            value = (raw_value or "").strip()
            row_dict[key] = value or None
        parsed_rows.append(row_dict)

    if not parsed_rows:
        raise ValidationError("The uploaded file has no data rows.", code="BIN_UPLOAD_NO_DATA_ROWS")
    if len(parsed_rows) > _MAX_ROWS:
        raise ValidationError(
            f"The uploaded file has {len(parsed_rows)} data rows — at most {_MAX_ROWS} are supported per upload.",
            code="BIN_UPLOAD_TOO_MANY_ROWS",
        )

    return parsed_rows


def parse_gift_card_upload_file(filename: str, content: bytes) -> List[GiftCardBulkUploadRowRequest]:
    """Backs POST .../gift-card/bulk-upload. See this module's docstring
    for the parsing rules; per-row business validation is unchanged,
    still owned entirely by
    app.services.gift_card_bin_service.bulk_upload_gift_card_bin_records."""
    table = _read_table(filename, content)
    rows = _rows_from_table(table, _GIFT_CARD_COLUMNS)
    return [GiftCardBulkUploadRowRequest(**row) for row in rows]


def parse_wallet_upload_file(filename: str, content: bytes) -> List[WalletBulkUploadRowRequest]:
    """Backs POST .../wallet/bulk-upload. See
    parse_gift_card_upload_file()'s docstring — identical reasoning."""
    table = _read_table(filename, content)
    rows = _rows_from_table(table, _WALLET_COLUMNS)
    return [WalletBulkUploadRowRequest(**row) for row in rows]
